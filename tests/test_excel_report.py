"""Tests du générateur de rapport Excel d'audit (cahier des charges §9)."""

from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from config.matching_config import MatchLevel, MatchStatus
from src.analysis import generer_synthese_mensuelle
from src.matching.appariement import Appariement, ResultatRapprochement
from src.models import LigneJournal, Periode, TransactionOM
from src.reports.excel_report import ExcelReportGenerator

JOUR = date(2026, 4, 16)
PERIODE = Periode.journee(JOUR)


def creer_jeu_test() -> ResultatRapprochement:
    """Crée un résultat de rapprochement riche avec conformes, anomalies, recette et doublons."""
    # Arrhes
    l_conforme = LigneJournal(
        ligne_source=2,
        date_operation=datetime(2026, 4, 16, 13, 0),
        client="ALPHA Jean",
        mode_paiement="Orange Money",
        montant=Decimal("90200"),
        reference_interne="RES-001",
    )
    l_ecart = LigneJournal(
        ligne_source=3,
        date_operation=datetime(2026, 4, 16, 14, 0),
        client="BETA Paul",
        mode_paiement="Orange Money",
        montant=Decimal("50000"),
        reference_interne="RES-002",
    )
    l_manquante = LigneJournal(
        ligne_source=4,
        date_operation=datetime(2026, 4, 16, 15, 0),
        client="GAMMA Luc",
        mode_paiement="Orange Money",
        montant=Decimal("30000"),
        reference_interne="RES-003",
    )
    l_doublon = LigneJournal(
        ligne_source=5,
        date_operation=datetime(2026, 4, 16, 16, 0),
        client="DELTA Marc",
        mode_paiement="Orange Money",
        montant=Decimal("20000"),
        reference_interne="RES-004",
    )

    # OM
    t_conforme = TransactionOM(
        numero=1,
        date_operation=JOUR,
        heure=time(13, 11),
        reference="MP260416.1311.A001",
        service="Merchant Payment",
        statut="Succès",
        compte_agent="656009773",
        correspondant="690000001",
        credit=Decimal("90200"),
        commission=Decimal("900"),
    )
    t_ecart = TransactionOM(
        numero=2,
        date_operation=JOUR,
        heure=time(14, 5),
        reference="MP260416.1405.A002",
        service="Merchant Payment",
        statut="Succès",
        compte_agent="656009773",
        correspondant="690000002",
        credit=Decimal("48000"),  # Écart de 2 000 FCFA
        commission=Decimal("480"),
    )
    t_recette = TransactionOM(
        numero=3,
        date_operation=JOUR,
        heure=time(15, 20),
        reference="MP260416.1520.A003",
        service="Merchant Payment",
        statut="Succès",
        compte_agent="691829711",
        correspondant="690000003",
        credit=Decimal("18000"),
        commission=Decimal("180"),
    )
    t_doublon = TransactionOM(
        numero=4,
        date_operation=JOUR,
        heure=time(16, 10),
        reference="MP260416.1610.A004",
        service="Merchant Payment",
        statut="Succès",
        compte_agent="696948928",
        correspondant="690000004",
        credit=Decimal("20000"),
        commission=Decimal("200"),
    )

    app_conforme = Appariement(
        lignes=(l_conforme,),
        transactions=(t_conforme,),
        niveau=MatchLevel.DATE_MONTANT,
        score=100.0,
    )
    app_ecart = Appariement(
        lignes=(l_ecart,),
        transactions=(t_ecart,),
        niveau=MatchLevel.DATE_MONTANT,
        score=100.0,
    )

    return ResultatRapprochement(
        periode=PERIODE,
        appariements=[app_conforme, app_ecart],
        arrhes_sans_om=[l_manquante],
        recette_du_jour=[t_recette],
        doublons_journal=[l_doublon],
        doublons_om=[t_doublon],
    )


def test_generation_rapport_excel_7_feuilles(tmp_path: Path):
    """Vérifie la génération complète du classeur et l'existence exacte des 7 feuilles du §9."""
    resultat = creer_jeu_test()
    generateur = ExcelReportGenerator(output_dir=tmp_path)
    fichier = generateur.generate(resultat, filename="test_rapport.xlsx")

    assert fichier.exists()
    assert fichier.stat().st_size > 0

    wb = openpyxl.load_workbook(fichier)
    feuilles_attendues = [
        "Synthèse",
        "Rapprochement",
        "Anomalies",
        "Manquants_Journal",
        "Manquants_OM",
        "Doublons",
        "Contrôle_Mensuel",
    ]
    assert wb.sheetnames == feuilles_attendues

    # Vérification feuille Synthèse
    ws_syn = wb["Synthèse"]
    assert "RAPPORT D'AUDIT" in ws_syn["A1"].value
    assert "Période auditée" in ws_syn["A2"].value
    assert "Statut Global" in ws_syn["A5"].value
    # Vérifier présence des contrôles
    labels_ctrl = [ws_syn.cell(row=r, column=1).value for r in range(8, 25)]
    assert "Nombre de lignes du journal" in labels_ctrl
    assert "Total Orange Money" in labels_ctrl
    assert "Montant des Écarts" in labels_ctrl

    # Vérifier Section 3 détaillée des opérations bloquantes
    textes_syn = [str(ws_syn.cell(row=r, column=1).value) for r in range(1, ws_syn.max_row + 1)]
    assert any("3. Détail des Opérations Bloquantes" in t for t in textes_syn)

    # Vérification feuille Rapprochement
    ws_rap = wb["Rapprochement"]
    assert ws_rap["A1"].value == "Date Opération"
    assert ws_rap["D1"].value == "Montant Journal"
    assert ws_rap["G1"].value == "Statut"
    assert ws_rap.auto_filter.ref is not None

    # Vérification feuille Anomalies
    ws_ano = wb["Anomalies"]
    assert ws_ano["A1"].value == "Date"
    assert ws_ano["H1"].value == "Bloquant Export"
    # L'écart de montant et l'arrhe manquante doivent être listés
    clients_ano = [ws_ano.cell(row=r, column=3).value for r in range(2, 5)]
    assert "BETA Paul" in clients_ano or "GAMMA Luc" in clients_ano
    bloquants_ano = [ws_ano.cell(row=r, column=8).value for r in range(2, 5)]
    assert "OUI" in bloquants_ano

    # Vérification feuille Contrôle_Mensuel
    ws_mens = wb["Contrôle_Mensuel"]
    assert ws_mens["A1"].value == "Date"
    assert ws_mens.cell(row=2, column=1).value == "16/04/2026"


def test_generation_rapport_en_memoire():
    """Vérifie la génération sous forme de flux binaire (BytesIO) pour Streamlit."""
    resultat = creer_jeu_test()
    generateur = ExcelReportGenerator()
    tampon = generateur.generate_bytes(resultat)

    assert tampon is not None
    assert tampon.getvalue()
    wb = openpyxl.load_workbook(tampon)
    assert len(wb.sheetnames) == 7
    assert "Synthèse" in wb.sheetnames


def test_generation_rapport_mensuel_campagne(tmp_path: Path):
    """Vérifie que le rapport Excel supporte un ResultatMensuel avec journées non couvertes."""
    from src.matching.campagne import JourneeNonCouverte, ResultatMensuel

    res1 = creer_jeu_test()
    non_couverte = JourneeNonCouverte(
        jour=date(2026, 4, 17),
        nb_transactions=2,
        total=Decimal("50000"),
    )
    mensuel = ResultatMensuel(
        periode=res1.periode,
        journees=[res1],
        non_couvertes=[non_couverte],
    )

    generateur = ExcelReportGenerator(output_dir=tmp_path)
    fichier = generateur.generate(mensuel, filename="test_rapport_mensuel.xlsx")
    assert fichier.exists()

    wb = openpyxl.load_workbook(fichier)
    ws_mens = wb["Contrôle_Mensuel"]
    statuts = [ws_mens.cell(row=r, column=9).value for r in range(2, 4)]
    assert "NON COUVERTE (SANS JOURNAL)" in statuts


def test_generation_rapprochement_colore_8_colonnes():
    """Vérifie le classeur de rapprochement dédié : 8 colonnes (avec Date et Opérateur) et surbrillances."""
    from src.analysis.reconciliation import LigneRapprochement
    from src.reports import generer_excel_rapprochement_colore

    l_conforme = LigneRapprochement(
        date_operation=date(2026, 4, 16),
        reference="MP260416.0001",
        client="ALPHA Jean",
        montant_journal=Decimal("90200"),
        montant_om=Decimal("90200"),
        ecart=Decimal("0"),
        statut=MatchStatus.CONFORME,
        observation="Rapprochement exact",
        operateur="Orange Money",
    )
    l_manquant = LigneRapprochement(
        date_operation=date(2026, 4, 16),
        reference="RES-003",
        client="GAMMA Luc",
        montant_journal=Decimal("30000"),
        montant_om=Decimal("0"),
        ecart=Decimal("30000"),
        statut=MatchStatus.MANQUANT_OM,
        observation="Arrhe sans encaissement Orange Money",
        operateur="Orange Money",
    )
    l_ecart = LigneRapprochement(
        date_operation=date(2026, 4, 16),
        reference="MP260416.0002",
        client="BETA Paul",
        montant_journal=Decimal("50000"),
        montant_om=Decimal("48000"),
        ecart=Decimal("2000"),
        statut=MatchStatus.ECART_MONTANT,
        observation="Écart de 2 000 FCFA",
        operateur="Orange Money",
    )
    l_croisement = LigneRapprochement(
        date_operation=date(2026, 4, 16),
        reference="MOMO-001",
        client="DELTA Marc",
        montant_journal=Decimal("20000"),
        montant_om=Decimal("20000"),
        ecart=Decimal("0"),
        statut=MatchStatus.CORRESPONDANCE_PROBABLE,
        observation="Croisement d'opérateur constaté",
        operateur="Orange Money ➔ MTN Mobile Money",
    )

    tampon = generer_excel_rapprochement_colore([l_conforme, l_manquant, l_ecart, l_croisement])
    wb = openpyxl.load_workbook(tampon)
    assert "Rapprochement" in wb.sheetnames
    ws = wb["Rapprochement"]

    # 1. Vérifier les 8 colonnes exactes
    colonnes = [ws.cell(row=1, column=c).value for c in range(1, 9)]
    assert colonnes == [
        "Date",
        "Client",
        "Opérateur",
        "Référence OM/MoMo",
        "Montant journal",
        "Montant Relevé",
        "Statut",
        "Observation",
    ]

    # 2. Ligne 2 : Conforme -> date formatée et toute la ligne en vert (DCFCE7)
    assert ws.cell(row=2, column=1).value == "16/04/2026"
    for c in range(1, 9):
        fill_color = ws.cell(row=2, column=c).fill.start_color.rgb
        assert fill_color in ("00DCFCE7", "DCFCE7")

    # 3. Ligne 3 : Manquant OM (non retrouvé) -> toute la ligne en rouge (FEE2E2)
    assert ws.cell(row=3, column=1).value == "16/04/2026"
    for c in range(1, 9):
        fill_color = ws.cell(row=3, column=c).fill.start_color.rgb
        assert fill_color in ("00FEE2E2", "FEE2E2")

    # 4. Ligne 4 : Écart montant -> cellules problématiques en jaune (FEF08A)
    # Montant journal (col 5), Montant Relevé (col 6), Statut (col 7)
    assert ws.cell(row=4, column=5).fill.start_color.rgb in ("00FEF08A", "FEF08A")
    assert ws.cell(row=4, column=6).fill.start_color.rgb in ("00FEF08A", "FEF08A")
    assert ws.cell(row=4, column=7).fill.start_color.rgb in ("00FEF08A", "FEF08A")
    # Date (col 1), Client (col 2), Opérateur (col 3) ont le fond neutre d'alerte (FFFBEB)
    assert ws.cell(row=4, column=1).fill.start_color.rgb in ("00FFFBEB", "FFFBEB")
    assert ws.cell(row=4, column=2).fill.start_color.rgb in ("00FFFBEB", "FFFBEB")

    # 5. Ligne 5 : Croisement opérateur -> Opérateur (col 3) et Statut (col 7) en jaune
    assert ws.cell(row=5, column=3).fill.start_color.rgb in ("00FEF08A", "FEF08A")
    assert ws.cell(row=5, column=7).fill.start_color.rgb in ("00FEF08A", "FEF08A")

    # 6. Ligne 6 : Totaux
    assert ws.cell(row=6, column=1).value == "TOTAL"
    assert ws.cell(row=6, column=5).value == "=SUM(E2:E5)"
    assert ws.cell(row=6, column=6).value == "=SUM(F2:F5)"


def test_section3_synthese_operations_bloquantes():
    """Vérifie que la Section 3 de la feuille Synthèse liste précisément les opérations bloquantes."""
    from src.analysis.reconciliation import LigneRapprochement
    from src.analysis.monthly_summary import generer_synthese_mensuelle

    res = creer_jeu_test()
    generateur = ExcelReportGenerator()
    tampon = generateur.generate_bytes(res)
    wb = openpyxl.load_workbook(tampon)
    ws_syn = wb["Synthèse"]

    # Trouver la ligne de la section 3
    row_sec3 = None
    for r in range(1, ws_syn.max_row + 1):
        val = ws_syn.cell(row=r, column=1).value
        if val and "3. Détail des Opérations Bloquantes" in str(val):
            row_sec3 = r
            break
    assert row_sec3 is not None, "Section 3 non trouvée dans la feuille Synthèse"

    # Vérifier l'en-tête de la section 3 (ligne suivante)
    row_header = row_sec3 + 1
    headers_sec3 = [ws_syn.cell(row=row_header, column=c).value for c in range(1, 9)]
    assert headers_sec3 == [
        "Date",
        "Client / Correspondant",
        "Référence",
        "Montant Journal",
        "Montant OM",
        "Écart",
        "Statut Bloquant",
        "Motif du Blocage & Action Requise",
    ]

    # Vérifier qu'au moins une opération bloquante y figure avec date, client, montant et action requise
    row_first_data = row_header + 1
    date_val = ws_syn.cell(row=row_first_data, column=1).value
    client_val = ws_syn.cell(row=row_first_data, column=2).value
    montant_val = ws_syn.cell(row=row_first_data, column=4).value
    action_val = ws_syn.cell(row=row_first_data, column=8).value

    assert date_val == "16/04/2026"
    assert client_val in ("BETA Paul", "GAMMA Luc")
    assert montant_val > 0
    assert "Action :" in str(action_val)


def test_section3_synthese_aucune_bloquante():
    """Vérifie que la Section 3 affiche un message de succès lorsque l'export est débloqué."""
    from src.models import Periode
    from src.analysis.monthly_summary import generer_synthese_mensuelle

    app_conforme = Appariement(
        lignes=(
            LigneJournal(
                ligne_source=2,
                date_operation=datetime(2026, 4, 16, 12, 0),
                client="ALPHA Jean",
                mode_paiement="Orange Money",
                montant=Decimal("50000"),
                reference_interne="RES-001",
            ),
        ),
        transactions=(
            TransactionOM(
                numero=1,
                date_operation=date(2026, 4, 16),
                heure=time(13, 11),
                reference="MP260416.0001",
                service="Merchant Payment",
                statut="Succès",
                compte_agent="698186110",
                correspondant="690000001",
                credit=Decimal("50000"),
            ),
        ),
        niveau=MatchLevel.DATE_MONTANT,
        score=100.0,
    )
    res_parfait = ResultatRapprochement(
        periode=Periode(date(2026, 4, 16), date(2026, 4, 16)),
        appariements=[app_conforme],
        recette_du_jour=[],
        arrhes_sans_om=[],
    )

    generateur = ExcelReportGenerator()
    tampon = generateur.generate_bytes(res_parfait)
    wb = openpyxl.load_workbook(tampon)
    ws_syn = wb["Synthèse"]

    # Trouver la section 3
    row_sec3 = None
    for r in range(1, ws_syn.max_row + 1):
        val = ws_syn.cell(row=r, column=1).value
        if val and "3. Détail des Opérations Bloquantes" in str(val):
            row_sec3 = r
            break
    assert row_sec3 is not None

    msg_succes = ws_syn.cell(row=row_sec3 + 1, column=1).value
    assert "Aucune opération bloquante détectée" in str(msg_succes)



def test_rapport_rapprochement_paiement_posterieur_ligne_entiere_jaune():
    """Un paiement postérieur à la saisie doit être entièrement surligné en jaune."""
    from src.analysis.reconciliation import LigneRapprochement
    from src.reports import generer_excel_rapprochement_colore

    ligne = LigneRapprochement(
        date_operation=date(2026, 6, 28),
        reference="MP260630.1310.A80050",
        client="Client BAL-8233",
        montant_journal=Decimal("50000"),
        montant_om=Decimal("50000"),
        ecart=Decimal("0"),
        statut=MatchStatus.PAIEMENT_POSTERIEUR_A_SAISIE,
        observation="Paiement effectué le 30/06/2026, soit 2 jour(s) après la saisie du 28/06/2026.",
        operateur="Orange Money",
    )

    tampon = generer_excel_rapprochement_colore([ligne])
    wb = openpyxl.load_workbook(tampon)
    ws = wb["Rapprochement"]

    assert ws.max_row == 3
    assert ws.cell(row=2, column=1).value == "28/06/2026"
    assert ws.cell(row=2, column=7).value == MatchStatus.PAIEMENT_POSTERIEUR_A_SAISIE.value

    for c in range(1, 9):
        fill_color = ws.cell(row=2, column=c).fill.start_color.rgb
        assert fill_color in ("00FEF08A", "FEF08A")


def test_rapport_rapprochement_paiement_posterieur_est_bloquant():
    """Le statut postérieur doit être visible comme anomalie bloquante dans le rapport."""
    from src.analysis.reconciliation import LigneRapprochement
    from src.reports import generer_excel_rapprochement_colore

    ligne = LigneRapprochement(
        date_operation=date(2026, 6, 19),
        reference="MP260622.1200.A7500",
        client="Client BAL-8089",
        montant_journal=Decimal("7500"),
        montant_om=Decimal("7500"),
        ecart=Decimal("0"),
        statut=MatchStatus.PAIEMENT_POSTERIEUR_A_SAISIE,
        observation="Paiement postérieur à la saisie : 22/06/2026 après 19/06/2026.",
        operateur="Orange Money",
    )

    tampon = generer_excel_rapprochement_colore([ligne])
    wb = openpyxl.load_workbook(tampon)
    ws = wb["Rapprochement"]

    assert ws.cell(row=2, column=7).value == "PAIEMENT_POSTERIEUR_A_SAISIE"
    assert "postérieur" in ws.cell(row=2, column=8).value.lower()
