"""Tests du lecteur de journal des encaissements.

C'est la source du rapprochement : elle porte tous les encaissements de la journée,
ventilés par mode de paiement. La colonne « Orange Money » délimite le périmètre.
"""

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest

from src.models import Periode
from src.readers import EncaissementsReader

FIXTURES = Path(__file__).resolve().parent / "fixtures"
DOUZE_MAI = FIXTURES / "encaissements_12-05-2026.xlsx"
DIX_HUIT_MAI = FIXTURES / "encaissements_18-05-2026.xlsx"


@pytest.fixture(scope="module")
def journal():
    return EncaissementsReader(DOUZE_MAI).read()


def test_la_periode_est_lue_dans_l_entete(journal):
    assert journal.periode == Periode.journee(date(2026, 5, 12))
    assert journal.periode.est_journee
    assert not journal.periode_deduite


def test_seule_la_colonne_orange_money_entre_dans_le_perimetre(journal):
    """Une ligne réglée en espèces est hors périmètre, pas une anomalie."""
    assert len(journal.mouvements) == 6
    montants = sorted(m.montant_om for m in journal.mouvements)
    assert montants == [
        Decimal("1500"), Decimal("2000"), Decimal("12000"),
        Decimal("75000"), Decimal("75000"), Decimal("120500"),
    ]


def test_la_nature_se_lit_dans_le_libelle(journal):
    """Arrhe ou facture : c'est le libellé qui tranche, pas le rapprochement.

    Cette distinction commande la nature de l'écriture comptable — une arrhe donne
    une ligne individuelle, une facture alimente la recette agrégée du jour.
    """
    arrhes = [m for m in journal.mouvements if m.est_arrhe]
    factures = [m for m in journal.mouvements if not m.est_arrhe]

    assert len(arrhes) == 3
    assert len(factures) == 3
    assert journal.total_arrhes == Decimal("270500")
    assert journal.total_factures == Decimal("15500")
    assert all(m.nature == "Arrhe" for m in arrhes)


def test_les_references_sont_extraites(journal):
    references = {m.reference for m in journal.mouvements}
    assert "9001" in references          # Réservation N°9001
    assert "KOT-2001" in references      # facture Kotibé
    assert "BAL-3001" in references      # facture Baleng


def test_le_fichier_annonce_ses_propres_totaux(journal):
    """Le récapitulatif sert de contrôle de lecture : un écart signale un oubli."""
    assert journal.total_declare_om == Decimal("286000")
    assert journal.total_declare_arrhes == Decimal("270500")
    assert journal.total_declare_factures == Decimal("15500")
    assert journal.total_om == journal.total_declare_om
    assert journal.ecart_au_recapitulatif() == 0


def test_le_recapitulatif_est_ecarte_de_la_table(journal):
    motifs = {motif for _, motif in journal.lot.rejets}
    assert motifs == {"bloc Récapitulatif"}


def test_une_journee_simple(journal):
    autre = EncaissementsReader(DIX_HUIT_MAI).read()
    assert autre.periode == Periode.journee(date(2026, 5, 18))
    assert autre.total_om == Decimal("9000")
    assert autre.ecart_au_recapitulatif() == 0


def test_les_lignes_canoniques_alimentent_le_moteur(journal):
    """Le moteur ne connaît que le modèle canonique, pas le format du journal."""
    lignes = journal.lignes_orange_money
    assert len(lignes) == 6
    assert all(ligne.est_orange_money for ligne in lignes)
    assert all(ligne.jour == date(2026, 5, 12) for ligne in lignes)
    assert sum(ligne.montant for ligne in lignes) == Decimal("286000")
    assert sum(1 for l in lignes if l.est_arrhe) == 3
    assert sum(1 for l in lignes if not l.est_arrhe) == 3


def test_fichier_absent():
    with pytest.raises(FileNotFoundError):
        EncaissementsReader(FIXTURES / "inexistant.xlsx").read()


def test_un_fichier_sans_colonne_orange_money(tmp_path):
    from openpyxl import Workbook

    chemin = tmp_path / "autre.xlsx"
    classeur = Workbook()
    classeur.active.append(["Mouvement", "Total", "Espèces"])
    classeur.save(chemin)
    with pytest.raises(ValueError, match="Orange Money"):
        EncaissementsReader(chemin).read()


def test_un_journal_sans_periode(tmp_path):
    """Sans période, le périmètre du contrôle est indéterminé : il faut refuser."""
    from openpyxl import Workbook

    chemin = tmp_path / "sans_periode.xlsx"
    classeur = Workbook()
    feuille = classeur.active
    feuille["B1"] = "Journal des encaissements"
    feuille.append([])
    for colonne, valeur in enumerate(["Mouvement", "Total", "Orange Money"], start=1):
        feuille.cell(row=2, column=colonne, value=valeur)
    classeur.save(chemin)
    with pytest.raises(ValueError, match="Période introuvable"):
        EncaissementsReader(chemin).read()


def test_lecture_journal_consolide_mensuel_avec_colonne_date(tmp_path):
    """Un journal consolidé mensuel portant une colonne Date doit affecter la bonne date à chaque mouvement."""
    from openpyxl import Workbook

    chemin = tmp_path / "encaissements_mensuel_mai_2026.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["B1"] = "Journal des encaissements\nPériode du 01/05/2026 au 31/05/2026 Facturé et encaissé\nHOTEL EXEMPLE SA"
    ws.append([])
    ws.append(["Date", "Mouvement", "Total", "Orange Money"])
    ws.append(["10/05/2026", "H-1001 50 000 FCFA # 101 DUPONT", 50000, 50000])
    ws.append(["12/05/2026", "KOT-2001 20 000 FCFA T 1", 20000, 20000])
    ws.append(["12/05/2026", "Réservation N°9001 13/05/26 ALPHA Arrhes", 75000, 75000])
    ws.append(["TOTAL PERIODE", None, 145000, 145000])
    wb.save(chemin)

    lecture = EncaissementsReader(chemin).read()
    assert lecture.periode == Periode(date(2026, 5, 1), date(2026, 5, 31))
    assert len(lecture.mouvements) == 3
    assert lecture.mouvements[0].date_mouvement == date(2026, 5, 10)
    assert lecture.mouvements[1].date_mouvement == date(2026, 5, 12)
    assert lecture.mouvements[2].date_mouvement == date(2026, 5, 12)

    # Ventilation par jour
    ventiles = lecture.ventiler_par_jour()
    assert len(ventiles) == 2
    assert ventiles[0].periode == Periode.journee(date(2026, 5, 10))
    assert len(ventiles[0].mouvements) == 1
    assert ventiles[0].total_om == Decimal("50000")

    assert ventiles[1].periode == Periode.journee(date(2026, 5, 12))
    assert len(ventiles[1].mouvements) == 2
    assert ventiles[1].total_om == Decimal("95000")
    assert all(l.jour == date(2026, 5, 12) for l in ventiles[1].lignes_orange_money)


def test_lecture_journal_consolide_mensuel_avec_sections(tmp_path):
    """Un journal consolidé avec séparateurs de sections par jour."""
    from openpyxl import Workbook

    chemin = tmp_path / "encaissements_sections_mai_2026.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["B1"] = "Journal des encaissements\nPériode du 01/05/2026 au 31/05/2026\nHOTEL EXEMPLE SA"
    ws.append([])
    ws.append(["Mouvement", "Total", "Orange Money"])
    ws.append(["Journée du 05/05/2026", None, None])
    ws.append(["BAL-1001 10 000 FCFA MB 1", 10000, 10000])
    ws.append(["Journée du 25/05/2026", None, None])
    ws.append(["BAL-2001 15 000 FCFA MB 2", 15000, 15000])
    wb.save(chemin)

    lecture = EncaissementsReader(chemin).read()
    assert len(lecture.mouvements) == 2
    assert lecture.mouvements[0].date_mouvement == date(2026, 5, 5)
    assert lecture.mouvements[1].date_mouvement == date(2026, 5, 25)

    ventiles = lecture.ventiler_par_jour()
    assert len(ventiles) == 2
    assert ventiles[0].periode.debut == date(2026, 5, 5)
    assert ventiles[1].periode.debut == date(2026, 5, 25)


def test_regularisations_annulees(tmp_path):
    """Les régularisations internes (+ et - pour la même arrhe) doivent s'annuler."""
    from openpyxl import Workbook

    chemin = tmp_path / "encaissements_regul.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Journal des encaissements\nPériode du 12/04/2026 au 12/04/2026"
    ws.append([])
    ws.append(["Mouvement", "Total", "Orange Money"])
    ws.append(["Facture H-7001 50 000 FCFA", 50000, 50000])
    ws.append(["Réservation N°8015 10/04/26 Régul Arrhes", -150000, -150000])
    ws.append(["Réservation N°8015 10/04/26 Arrhes Régul", 150000, 150000])
    ws.append(["TOTAL PERIODE", None, 50000])
    wb.save(chemin)

    lecture = EncaissementsReader(chemin, nom_source="mon_journal.xlsx").read()
    assert lecture.source == "mon_journal.xlsx"
    assert len(lecture.mouvements) == 1
    assert lecture.mouvements[0].montant_om == Decimal("50000")
    assert lecture.total_om == Decimal("50000")
    assert lecture.ecart_au_recapitulatif() == Decimal("0")


