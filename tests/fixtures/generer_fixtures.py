"""Génère les classeurs d'exemple utilisés par les tests de lecture.

Les fichiers réels contiennent des noms de clients et des numéros de téléphone : ils
restent dans `data/input/`, exclu du dépôt. Ces exemples en reproduisent **la
structure et les pièges** — bloc d'en-tête portant la période, table décalée,
ventilation par mode de paiement, récapitulatif auto-vérifiant, sous-relevés
concaténés, en-têtes brouillés par les cellules fusionnées — avec des données
entièrement fabriquées.

Régénération :

    .venv/bin/python tests/fixtures/generer_fixtures.py
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

ICI = Path(__file__).resolve().parent


def _entetes_releve(ws, brouilles: bool) -> None:
    """Ajoute la paire ligne de groupe / ligne d'en-tête.

    Dans le fichier réel, seule la première paire est exploitable : les cellules
    fusionnées décalent les suivantes au point d'y écrire « Généré le : » en guise de
    nom de colonne. Le lecteur doit donc s'en tenir à la première.
    """
    if brouilles:
        ws.append([None, None, None, None, "Type de rapport :", None, None, None,
                   "Coordonnées entrepris", None, None, "Date", None, "Montant (XAF)",
                   None, "Commissions (XAF)"])
        ws.append(["Agent", "Opération", "Paiement", "N° de Compte", "Statut", "Réseau :",
                   "Compte Orange Money :", "Contact chez le client :", "Généré le :",
                   "Heure", "Mode", "Généré le :", "Heure", "N°", "Référence",
                   "Compte: 600000001", "Sous-réseau"])
        return
    ws.append([None, None, None, None, "Opération", None, None, None, "Agent", None, None,
               "Correspondant", None, "Montant (XAF)", None, "Commissions (XAF)"])
    ws.append(["N°", "Date", "Heure", "Référence", "Service", "Paiement", "Statut", "Mode",
               "N° de Compte", "Wallet", "N° Pseudo", "N° de Compte", "Wallet",
               "Débit", "Crédit", "Compte: 600000001", "Sous-réseau"])


def _transaction(numero, date, heure, reference, service, statut, agent, correspondant,
                 debit=None, credit=None, commission=None):
    return [numero, date, heure, reference, service, "Transaction", statut, "USSD",
            agent, "Normal", None, correspondant, "Normal", debit, credit, 0, commission]


def releve_orange_money() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Channel User Transaction Report"

    ws.append(["."])
    ws.append(["RELEVE DES OPERATIONS OM DU MOIS DE MAI 2026"])
    ws.append(["Application :", None, None, "Orange Money"])
    ws.append(["Réseau :", None, None, "Cameroon"])
    ws.append(["Début de Période :", None, None, "01/05/2026"])
    ws.append(["Fin de Période :", None, None, "31/05/2026"])
    ws.append(["Type de rapport :", None, None, "Rapport mensuel"])
    ws.append(["."])

    # --- premier sous-relevé : le préambule utilise ses propres étiquettes
    ws.append(["Compte Orange Money :", None, None, "600000001"])
    ws.append(["Coordonnées entrepris", None, None, "POINT DE VENTE UN - Cameroun, 000000, Ville"])
    ws.append(["Contact chez le client :", None, None, "USSD - 600000001 RELEVE MAI 2026 ACCUEIL"])
    ws.append(["."])
    _entetes_releve(ws, brouilles=False)
    ws.append([None, None, None, "Transactions échouées", "Wallet Normal", None, None, None,
               "600000001", "Normal", None, None, None, None, 0])
    ws.append(_transaction(1, "03/05/2026", "10:00:00", "MP260503.1000.A00001",
                           "Merchant Payment", "Echec", "600000001", "690000011", credit=50000))
    ws.append(["Total", None, None, "Transactions échouées", "Wallet Normal", None, None, None,
               "600000001", "Normal", None, "Total échecs", None, None, 50000])
    ws.append([None, None, None, "Transactions réussies", "Wallet Normal", None, None, None,
               "600000001", "Normal", None, "Solde initial", None, None, 100000])
    ws.append(_transaction(2, "12/05/2026", "11:10:44", "MP260512.1110.B00002",
                           "Merchant Payment", "Succès", "600000001", "690000012",
                           credit=75000, commission=-750))
    ws.append(_transaction(3, "12/05/2026", "14:02:19", "MP260512.1402.C00003",
                           "Merchant Payment", "Succès", "600000001", "690000013",
                           credit=120500, commission=-1205))
    ws.append(_transaction(4, "12/05/2026", "17:35:08", "MP260512.1735.D00004",
                           "Merchant Payment", "Succès", "600000001", "690000014",
                           credit=75000, commission=-750))
    ws.append(_transaction(5, "12/05/2026", "20:11:02", "MP260512.2011.E00005",
                           "Merchant Payment", "Succès", "600000001", "690000015",
                           credit=3500, commission=-35))
    ws.append(_transaction(6, "18/05/2026", "08:45:00", "MP260518.0845.F00006",
                           "Merchant Payment", "Succès", "600000001", "690000016",
                           credit=9000, commission=-90))
    ws.append(["Relevé de vos opérations", None, None, "Transactions réussies", "Wallet Normal",
               None, None, None, "600000001", "Normal", None, "Total activités", None, None, 283000])
    ws.append(["."])

    # --- second sous-relevé : préambule ré-étiqueté, en-têtes brouillés
    ws.append(["Application :", None, None, "600000002"])
    ws.append(["Fin de Période :", None, None, "POINT DE VENTE DEUX - Cameroun, 000000, Ville"])
    ws.append(["Orange Money", None, None, "USSD - 600000002 RELEVE MAI 2026 RESTAURANT"])
    ws.append(["."])
    _entetes_releve(ws, brouilles=True)
    ws.append(_transaction(7, "12/05/2026", "19:20:30", "MP260512.1920.G00007",
                           "Merchant Payment", "Succès", "600000002", "690000017",
                           credit=12000, commission=-120))
    ws.append(_transaction(8, "12/05/2026", "21:05:57", "PP260512.2105.H00008",
                           "C2C Transfer", "Succès", "600000002", "600000001",
                           credit=200000))
    # Ligne de commission : un numéro, aucune date.
    ws.append([9, None, None, None, "  Commissions", None, None, None, "600000002", "Normal",
               None, "600000002", None, 4500])
    ws.append(_transaction(10, "20/05/2026", "10:15:00", "MP260520.1015.I00010",
                           "Merchant Payment", "Succès", "600000002", "690000018",
                           credit=45000, commission=-450))
    ws.append(["Relevé de vos opérations", None, None, "Transactions réussies", "Wallet Normal",
               None, None, None, "600000002", "Normal", None, "Solde final", None, None, 95000])
    return wb


MODES = ["Total", "Espèces", "Chèque", "Carte Bancaire", "Crédit", "Transfert PV",
         "Virement Bancaire", "Orange Money", "MTN Mobile Money", "CB Visa",
         "Arrhes/Acpt réintégrés"]
COL_OM = MODES.index("Orange Money")


def _mouvement(libelle, total, **modes):
    """Une ligne du journal : le libellé, le total, puis la ventilation par mode."""
    ligne = [libelle] + [None] * len(MODES)
    ligne[1] = total
    for mode, montant in modes.items():
        ligne[1 + MODES.index(mode.replace("_", " "))] = montant
    return ligne


def journal_des_encaissements(jour: str, mouvements, recapitulatif) -> Workbook:
    """Journal des encaissements d'une journée, ventilé par mode de paiement.

    Reproduit la structure réelle : bloc d'en-tête portant la période, table des
    mouvements, bloc RECAPITULATIF qui annonce ses propres totaux, note et pagination.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "A"

    ws["B1"] = (
        f"Journal des  encaissements \nPériode du {jour} au {jour} Facturé et encaissé\n"
        "HOTEL EXEMPLE SA\n"
    )
    ws.append([])
    for colonne, valeur in enumerate(["Mouvement"] + MODES, start=1):
        ws.cell(row=2, column=colonne, value=valeur)

    for ligne in mouvements:
        ws.append(ligne)

    ws.append(["RECAPITULATIF"] + MODES)
    for libelle, total, om in recapitulatif:
        ligne = [libelle] + [None] * len(MODES)
        ligne[1] = total
        ligne[1 + COL_OM] = om
        ws.append(ligne)
    ws.append([None, None, None, None, "Total facturé : exemple"])
    ws.append([None] * 18 + ["1/1"])
    return wb


def journal_12_mai() -> Workbook:
    """Journée qui se rapproche entièrement, avec un regroupement de deux factures."""
    mouvements = [
        _mouvement("H-1001 50 000 FCFA # 101 ENTREPRISE", 50000, Espèces=50000),
        _mouvement("KOT-2001 2 000 FCFA T 1", 2000, Orange_Money=2000),
        _mouvement("KOT-2002 1 500 FCFA T 2", 1500, Orange_Money=1500),
        _mouvement("BAL-3001 12 000 FCFA MB 4", 12000, Orange_Money=12000),
        _mouvement("Réservation N°9001 13/05/26  ALPHA Arrhes  M. ALPHA Jean", 75000, Orange_Money=75000),
        _mouvement("Réservation N°9002 13/05/26  BETA Arrhes  Mme BETA Claire", 120500, Orange_Money=120500),
        _mouvement("Réservation N°9003 14/05/26  GAMMA Arrhes  M. GAMMA Paul", 75000, Orange_Money=75000),
        _mouvement("Réservation N°9004 14/05/26  DELTA Arrhes  Mme DELTA Sophie", 40000, Virement_Bancaire=40000),
    ]
    recapitulatif = [
        ("Encaissement de factures", 65500, 15500),
        ("Encaissement d'arrhes", 310500, 270500),
        ("Total période *", 376000, 286000),
    ]
    return journal_des_encaissements("12/05/2026", mouvements, recapitulatif)


def journal_18_mai() -> Workbook:
    """Journée simple, un seul encaissement Orange Money."""
    mouvements = [
        _mouvement("BAL-3002 9 000 FCFA MB 1", 9000, Orange_Money=9000),
        _mouvement("H-1002 20 000 FCFA # 102 CLIENT", 20000, Espèces=20000),
    ]
    recapitulatif = [
        ("Encaissement de factures", 29000, 9000),
        ("Total période *", 29000, 9000),
    ]
    return journal_des_encaissements("18/05/2026", mouvements, recapitulatif)


if __name__ == "__main__":
    releve_orange_money().save(ICI / "releve_om_exemple.xlsx")
    journal_12_mai().save(ICI / "encaissements_12-05-2026.xlsx")
    journal_18_mai().save(ICI / "encaissements_18-05-2026.xlsx")
    print("fixtures générées dans", ICI)
