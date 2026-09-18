"""Styles, palettes de couleurs et utilitaires de mise en forme pour les rapports Excel (OpenPyXL)."""

from typing import Optional
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config.matching_config import MatchStatus

# --- Palette de couleurs ---
NAVY_HEADER = "1E3A8A"
NAVY_SUBHEADER = "2563EB"
ORANGE_OM = "FF6600"
WHITE = "FFFFFF"
GRAY_LIGHT = "F8FAFC"
GRAY_BORDER = "E2E8F0"
GRAY_DARK = "334155"

# Statuts
STATUS_COLORS: dict[MatchStatus, tuple[str, str]] = {
    MatchStatus.CONFORME: ("DCFCE7", "166534"),               # Vert doux
    MatchStatus.CORRESPONDANCE_PROBABLE: ("E0F2FE", "075985"),# Bleu clair
    MatchStatus.CORRESPONDANCE_GROUPEE: ("F3E8FF", "6B21A8"), # Violet doux
    MatchStatus.RECETTE_JOUR: ("F1F5F9", "475569"),           # Gris-ardoise doux
    MatchStatus.ECART_MONTANT: ("FEE2E2", "991B1B"),          # Rouge doux
    MatchStatus.MANQUANT_OM: ("FFE4E6", "9F1239"),            # Rose / rouge bordeaux
    MatchStatus.MANQUANT_JOURNAL: ("FFEDD5", "9A3412"),       # Orange doux
    MatchStatus.A_CONTROLER: ("FEF3C7", "92400E"),            # Ambre
    MatchStatus.STATUT_OM_INVALIDE: ("F3F4F6", "6B7280"),     # Gris
    MatchStatus.PAIEMENT_POSTERIEUR_A_SAISIE: ("FEF08A", "854D0E"),     # Jaune pâle
}

# --- Formats numériques ---
FORMAT_MONTANT_FCFA = '#,##0 "FCFA"'
FORMAT_NOMBRE_ENTIER = '#,##0'
FORMAT_POURCENTAGE = '0.0 "%"'
FORMAT_DATE = 'DD/MM/YYYY'

# --- Polices ---
FONT_TITLE = Font(name="Calibri", size=15, bold=True, color=NAVY_HEADER)
FONT_SECTION = Font(name="Calibri", size=12, bold=True, color=NAVY_HEADER)
FONT_HEADER = Font(name="Calibri", size=10, bold=True, color=WHITE)
FONT_DATA = Font(name="Calibri", size=10, color="0F172A")
FONT_DATA_BOLD = Font(name="Calibri", size=10, bold=True, color="0F172A")
FONT_MUTED = Font(name="Calibri", size=9, italic=True, color="64748B")

# --- Remplissages ---
FILL_HEADER = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
FILL_SUBHEADER = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
FILL_ZEBRA_EVEN = PatternFill(start_color=GRAY_LIGHT, end_color=GRAY_LIGHT, fill_type="solid")
FILL_ZEBRA_ODD = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")
FILL_TOTAL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

# --- Bordures ---
BORDER_THIN = Border(
    left=Side(style="thin", color=GRAY_BORDER),
    right=Side(style="thin", color=GRAY_BORDER),
    top=Side(style="thin", color=GRAY_BORDER),
    bottom=Side(style="thin", color=GRAY_BORDER),
)
BORDER_TOP_BOTTOM_DOUBLE = Border(
    top=Side(style="thin", color=GRAY_DARK),
    bottom=Side(style="double", color=GRAY_DARK),
)

# --- Alignements ---
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_HEADER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def appliquer_style_entete(cell, fill=FILL_HEADER, font=FONT_HEADER, alignment=ALIGN_HEADER):
    """Applique le style standard d'en-tête de colonne."""
    cell.fill = fill
    cell.font = font
    cell.alignment = alignment
    cell.border = BORDER_THIN


def appliquer_style_donnees(cell, is_even: bool = False, is_bold: bool = False, align=ALIGN_LEFT):
    """Applique le style de cellule de données avec zébrage."""
    cell.fill = FILL_ZEBRA_EVEN if is_even else FILL_ZEBRA_ODD
    cell.font = FONT_DATA_BOLD if is_bold else FONT_DATA
    cell.alignment = align
    cell.border = BORDER_THIN


def appliquer_style_statut(cell, statut: MatchStatus):
    """Colorise la cellule de statut avec sa teinte et police associées."""
    bg_color, fg_color = STATUS_COLORS.get(statut, ("FFFFFF", "000000"))
    cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    cell.font = Font(name="Calibri", size=10, bold=True, color=fg_color)
    cell.alignment = ALIGN_CENTER
    cell.border = BORDER_THIN


def ajuster_largeurs_colonnes(ws, min_width: int = 12, max_width: int = 50, padding: int = 3):
    """Calcule et applique la largeur idéale de chaque colonne selon son contenu."""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = cell.value
            if val is not None:
                txt = str(val)
                # Prendre la ligne la plus longue si retour à la ligne
                lignes = txt.split("\n")
                longueur = max(len(l) for l in lignes)
                if longueur > max_len:
                    max_len = longueur
        largeur = max(min(max_len + padding, max_width), min_width)
        ws.column_dimensions[col_letter].width = largeur
